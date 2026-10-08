import pathlib
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

D = pathlib.Path(__file__).parent
IMG = D / "img"

INK = RGBColor(0x14, 0x1B, 0x26)
BODY = RGBColor(0x47, 0x52, 0x61)
MUTE = RGBColor(0x7C, 0x87, 0x95)
BLUE = RGBColor(0x25, 0x63, 0xEB)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
EDGE = RGBColor(0xDF, 0xE4, 0xEA)
DARK = RGBColor(0x0F, 0x13, 0x19)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MONOBG = RGBColor(0xEE, 0xF2, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W = 13.333


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def tb(slide, x, y, w, h, text, size=16, bold=False, color=BODY, align=PP_ALIGN.LEFT,
       font="Segoe UI", space=4, anchor=MSO_ANCHOR.TOP, line=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text.split("\n") if isinstance(text, str) else text
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        if line:
            p.line_spacing = line
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return box


def rect(slide, x, y, w, h, fill=CARD, line=EDGE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = adj
        except Exception:
            pass
    s.text_frame.text = ""
    return s


def header(slide, kicker, title, sub=None):
    rect(slide, 0, 0, W, 0.09, fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(slide, 0.72, 0.46, 11.9, 0.3, kicker.upper(), 11, True, BLUE)
    tb(slide, 0.72, 0.76, 11.9, 0.62, title, 29, True, INK)
    if sub:
        tb(slide, 0.72, 1.44, 11.9, 0.45, sub, 14.5, False, MUTE)


def mono(slide, x, y, w, h, text, size=11.5, color=INK, bg=MONOBG):
    if bg is not None:
        rect(slide, x, y, w, h, fill=bg, line=None, adj=0.12)
    tb(slide, x + 0.16, y + 0.11, w - 0.32, h - 0.22, text, size, False, color,
       font="Consolas", space=3)


def footer(slide, n):
    tb(slide, 0.72, 7.02, 8, 0.25, "Self-hosted room booking for Cisco RoomOS  ·  concept", 9.5, False, MUTE)
    tb(slide, 11.3, 7.02, 1.3, 0.25, str(n), 9.5, False, MUTE, align=PP_ALIGN.RIGHT)


def chip(slide, x, y, w, h, label, color, fill):
    rect(slide, x, y, w, h, fill=fill, line=None, adj=0.35)
    tb(slide, x, y + h / 2 - 0.12, w, 0.3, label, 11.5, True, color, align=PP_ALIGN.CENTER)


N = [0]


def num():
    N[0] += 1
    return N[0]


# ---------------------------------------------------------------- 1 title
s = blank()
rect(s, 0, 0, W, 7.5, fill=DARK, line=None, shape=MSO_SHAPE.RECTANGLE)
rect(s, 0, 0, 0.16, 7.5, fill=GREEN, line=None, shape=MSO_SHAPE.RECTANGLE)
tb(s, 1.1, 2.15, 11, 0.4, "CONCEPT PROPOSAL", 13, True, RGBColor(0x4A, 0xDE, 0x80))
tb(s, 1.1, 2.68, 11.2, 1.5, "Self-hosted room booking\nfor Cisco RoomOS", 46, True, WHITE, line=1.05)
tb(s, 1.1, 4.5, 10.5, 0.9,
   "Keeping the Cisco panels and video devices. Removing the Webex cloud entirely.\n"
   "Booking data stays in Exchange Online, reached directly over Microsoft Graph.",
   17, False, RGBColor(0x9A, 0xA4, 0xB2), line=1.3)
chip(s, 1.1, 5.75, 2.5, 0.42, "Provisioning Mode: Off", RGBColor(0x4A, 0xDE, 0x80), RGBColor(0x10, 0x28, 0x1A))
chip(s, 3.75, 5.75, 2.3, 0.42, "Persistent Web App", RGBColor(0x93, 0xC5, 0xFD), RGBColor(0x12, 0x1E, 0x33))
chip(s, 6.2, 5.75, 2.1, 0.42, "Manual LED control", RGBColor(0xFB, 0xBF, 0x24), RGBColor(0x2A, 0x22, 0x08))
chip(s, 8.45, 5.75, 2.3, 0.42, "Microsoft Graph", RGBColor(0x93, 0xC5, 0xFD), RGBColor(0x12, 0x1E, 0x33))

# ---------------------------------------------------------------- 2 constraint
s = blank()
header(s, "Starting point", "The constraint, and what we are protecting",
       "Webex is no longer available to us. The hardware and the calendar are.")
cards = [
    ("Webex cloud is out", RED,
     "No Control Hub registration, no Hybrid Calendar,\nno Webex device management or analytics.\n\n"
     "This is a hard constraint, not a preference."),
    ("Cisco hardware stays", GREEN,
     "Room Navigator panels and RoomOS video devices\nare a significant investment and work well.\n\n"
     "We keep every device, including the LED strips."),
    ("Calendar stays in M365", BLUE,
     "Room mailboxes in Exchange Online remain the\nsystem of record. Users keep booking in Outlook.\n\n"
     "We reach them directly via Microsoft Graph."),
]
for i, (t, c, body) in enumerate(cards):
    x = 0.72 + i * 4.06
    rect(s, x, 2.25, 3.78, 2.9)
    rect(s, x, 2.25, 3.78, 0.075, fill=c, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(s, x + 0.3, 2.55, 3.2, 0.4, t, 17, True, INK)
    tb(s, x + 0.3, 3.08, 3.2, 1.9, body, 12.5, False, BODY, line=1.35)
rect(s, 0.72, 5.5, 11.89, 1.0, fill=RGBColor(0xF1, 0xF6, 0xFE), line=RGBColor(0xC7, 0xDC, 0xFC))
tb(s, 1.05, 5.72, 11.3, 0.6,
   "The question this concept answers:  can a Room Navigator still act as a room booking panel, with occupancy-driven "
   "room release, when it is never registered to Webex?", 14.5, True, RGBColor(0x1E, 0x40, 0xAF), line=1.3)
footer(s, num() + 1)

# ---------------------------------------------------------------- 3 device decision
s = blank()
header(s, "The device decision", "Three independent settings define a Room Navigator",
       "Each is set locally over the xAPI or the device web interface. No cloud involved.")
rows = [
    ("xCommand Provisioning SetType", "PairedToCodec  /  Standalone", "Standalone", "Physical topology — no codec behind the panel"),
    ("xCommand SystemUnit SetTouchPanelMode", "Controller / Scheduler / PersistentWebApp", "PersistentWebApp", "What the screen runs. Controller needs a codec"),
    ("xConfiguration Provisioning Mode", "Off  /  Auto  /  Webex", "Off", "Off = \u201ccustomer managed\u201d. This is our zero-cloud switch"),
]
y = 2.3
tb(s, 0.9, 2.02, 4.2, 0.25, "SETTING", 10, True, MUTE)
tb(s, 5.3, 2.02, 2.9, 0.25, "VALUES", 10, True, MUTE)
tb(s, 8.3, 2.02, 1.7, 0.25, "WE USE", 10, True, MUTE)
for i, (k, v, pick, why) in enumerate(rows):
    rect(s, 0.72, y, 9.3, 1.12)
    tb(s, 0.95, y + 0.18, 4.3, 0.3, k, 11.5, False, INK, font="Consolas")
    tb(s, 0.95, y + 0.55, 7.2, 0.35, why, 12, False, MUTE)
    tb(s, 5.3, y + 0.18, 2.9, 0.3, v, 11, False, BODY, font="Consolas")
    chip(s, 8.32, y + 0.14, 1.5, 0.38, pick, WHITE, BLUE if i < 2 else GREEN)
    y += 1.25
rect(s, 10.2, 2.3, 2.41, 3.62, fill=RGBColor(0x10, 0x28, 0x1A), line=None)
tb(s, 10.45, 2.6, 1.95, 2.6,
   "Result\n\nA wall panel we own\nend to end.\n\nNo Webex account.\nNo Control Hub.\nNo cloud dependency\nin the booking path.",
   13, True, RGBColor(0x4A, 0xDE, 0x80), line=1.45)
rect(s, 0.72, 6.1, 11.89, 0.82, fill=RGBColor(0xFF, 0xF8, 0xEB), line=RGBColor(0xF5, 0xD9, 0x9C))
tb(s, 1.05, 6.28, 11.3, 0.5,
   "Trade-off:  in customer-managed mode Cisco documents the entire Bookings xAPI family as not applicable. "
   "We therefore build the booking UI ourselves — and Cisco's own setup guide prescribes exactly that.",
   13, False, RGBColor(0x92, 0x5F, 0x08), line=1.25)
footer(s, num() + 1)

# ---------------------------------------------------------------- 4 what we build instead
s = blank()
header(s, "Consequence", "What customer-managed mode gives and takes",
       "Taken from the Cisco Room Navigator (stand-alone) API Reference Guide, D15512.02, RoomOS 11.9.")
rect(s, 0.72, 2.2, 5.85, 4.1, fill=RGBColor(0xFE, 0xF4, 0xF4), line=RGBColor(0xF3, 0xCF, 0xCF))
tb(s, 1.05, 2.42, 5.2, 0.35, "NOT AVAILABLE   ·   customer managed", 11, True, RED)
lost = ["xCommand Bookings Put / Book / Delete / Get / List",
        "xCommand Bookings Respond / NotificationSnooze",
        "xStatus Bookings Availability / Current Id",
        "xConfiguration RoomScheduler Enabled",
        "The native Cisco Room Scheduler screen",
        "LED auto-driven from booking state"]
yy = 2.92
for item in lost:
    tb(s, 1.28, yy, 5.1, 0.3, "×   " + item, 12.5, False, RGBColor(0x99, 0x1B, 0x1B))
    yy += 0.46
tb(s, 1.05, 5.72, 5.2, 0.45,
   "Each is marked \u201cDoesn't apply for a customer\nmanaged Room Navigator\u201d.", 11, False, RGBColor(0xB4, 0x5A, 0x5A), line=1.3)

rect(s, 6.76, 2.2, 5.85, 4.1, fill=RGBColor(0xF1, 0xF9, 0xF4), line=RGBColor(0xC6, 0xE6, 0xD2))
tb(s, 7.09, 2.42, 5.2, 0.35, "AVAILABLE   ·   and officially prescribed", 11, True, GREEN)
kept = ["Persistent Web App — full-screen, undismissable",
        "Local xAPI over WebSocket, HTTP and SSH",
        "UserInterface LedControl Mode: Manual",
        "LedControl Color Set — Green / Yellow / Red / Off",
        "Panel temperature and humidity sensors",
        "Full local admin, no cloud account"]
yy = 2.92
for item in kept:
    tb(s, 7.32, yy, 5.1, 0.3, "✓   " + item, 12.5, False, RGBColor(0x14, 0x6C, 0x43))
    yy += 0.46
tb(s, 7.09, 5.72, 5.2, 0.45,
   "Cisco's own customer-managed walkthrough includes\nsetting LedControl Mode to Manual.", 11, False, RGBColor(0x4E, 0x8A, 0x68), line=1.3)
footer(s, num() + 1)

# ---------------------------------------------------------------- 5 architecture
s = blank()
header(s, "Architecture", "One broker, two devices per room, one calendar",
       "Only the calendar lives outside. Everything else runs on your own infrastructure.")

rect(s, 0.72, 2.12, 4.05, 3.5, fill=RGBColor(0xFA, 0xFB, 0xFC), line=RGBColor(0xC9, 0xD1, 0xDA), lw=1.25)
tb(s, 0.95, 2.26, 4, 0.28, "MEETING ROOM", 10, True, MUTE)

# panel
rect(s, 0.95, 2.72, 3.58, 1.22, fill=DARK, line=None)
tb(s, 1.3, 2.9, 3.0, 0.3, "Room Navigator", 14, True, WHITE)
tb(s, 1.3, 3.23, 3.1, 0.55, "standalone · outside room\nPersistent Web App", 10.5, False,
   RGBColor(0x9A, 0xA4, 0xB2), line=1.25, space=1)
rect(s, 1.08, 2.8, 0.08, 1.06, fill=GREEN, line=None, shape=MSO_SHAPE.RECTANGLE)

# codec
rect(s, 0.95, 4.22, 3.58, 1.22, fill=RGBColor(0x1C, 0x22, 0x2C), line=None)
tb(s, 1.3, 4.4, 3.0, 0.3, "RoomOS video device", 14, True, WHITE)
tb(s, 1.3, 4.73, 3.1, 0.55, "in room · has its own Navigator\noccupancy source only", 10.5, False,
   RGBColor(0x9A, 0xA4, 0xB2), line=1.25, space=1)

# broker
rect(s, 5.55, 3.18, 2.6, 1.75, fill=BLUE, line=None)
tb(s, 5.7, 3.48, 2.3, 0.35, "Room Broker", 15, True, WHITE, align=PP_ALIGN.CENTER)
tb(s, 5.7, 3.9, 2.3, 0.8, "reconciler\noccupancy state machine\ncalendar drivers", 10, False, RGBColor(0xD6, 0xE4, 0xFF),
   align=PP_ALIGN.CENTER, line=1.3, space=1)

for y0, y1, lbl in ((3.33, 3.6, "UI + LED"), (4.83, 4.5, "sensors")):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(4.53), Inches(y0), Inches(5.55), Inches(y1))
    c.line.color.rgb = RGBColor(0x9C, 0xA8, 0xB6)
    c.line.width = Pt(1.5)
    tb(s, 4.55, (y0 - 0.3) if y0 < 4 else (y0 + 0.06), 1.0, 0.25, lbl, 9, False, MUTE, align=PP_ALIGN.CENTER)

# calendar
rect(s, 8.75, 2.12, 3.86, 3.5, fill=RGBColor(0xF6, 0xF8, 0xFB), line=EDGE)
tb(s, 9.0, 2.26, 3.3, 0.28, "CALENDAR SOURCE", 10, True, MUTE)
rect(s, 9.0, 2.72, 3.36, 1.0, fill=CARD, line=EDGE)
tb(s, 9.22, 2.9, 3.0, 0.3, "Exchange Online", 14, True, INK)
tb(s, 9.22, 3.22, 3.0, 0.3, "Microsoft Graph · app-only", 10.5, False, MUTE)
for i, (nm, note) in enumerate([("Exchange on-prem", "EWS · streaming"),
                                ("Google Workspace", "sync tokens"),
                                ("CalDAV", "Nextcloud / SOGo")]):
    yy = 3.9 + i * 0.57
    rect(s, 9.0, yy, 3.36, 0.47, fill=RGBColor(0xFB, 0xFC, 0xFD), line=EDGE)
    tb(s, 9.22, yy + 0.12, 1.9, 0.25, nm, 11, False, BODY)
    tb(s, 11.0, yy + 0.13, 1.3, 0.25, note, 9.5, False, MUTE, align=PP_ALIGN.RIGHT)
c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(8.15), Inches(4.05), Inches(8.75), Inches(3.5))
c.line.color.rgb = RGBColor(0x9C, 0xA8, 0xB6)
c.line.width = Pt(1.5)
tb(s, 8.0, 4.15, 0.95, 0.25, "Graph", 9, False, MUTE, align=PP_ALIGN.CENTER)

rect(s, 0.72, 5.95, 11.89, 0.75, fill=RGBColor(0xF1, 0xF6, 0xFE), line=RGBColor(0xC7, 0xDC, 0xFC))
tb(s, 1.05, 6.12, 11.3, 0.45,
   "The panel and the video device never talk to each other. The broker is what joins them — "
   "replacing the Workspace association that Webex used to provide.", 13, False, RGBColor(0x1E, 0x40, 0xAF), line=1.25)
footer(s, num() + 1)

# ---------------------------------------------------------------- 6-8 panel images
panels = [
    ("panel_available.png", "The panel · available",
     "Room free. LED green. Tapping a duration books the room instantly.",
     [("Ad-hoc booking", "15 / 30 / 60 min, or a custom slot"),
      ("Today's agenda", "read from the room mailbox via Graph"),
      ("Panel sensors", "temperature and humidity come from the Navigator itself"),
      ("LED strip", "Green — driven by the broker, not by RoomOS")]),
    ("panel_inuse.png", "The panel · in use",
     "Meeting in progress and occupancy confirmed. LED red.",
     [("Live occupancy", "people count streamed from the in-room video device"),
      ("Extend", "lengthens the booking if the next slot is free"),
      ("End & release", "frees the room immediately for everyone"),
      ("LED strip", "Red — set by the broker on booking start")]),
    ("panel_checkin.png", "The panel · check-in and auto-release",
     "Booked but empty. The room is reclaimed if nobody shows up.",
     [("Check-in window", "configurable, default 10 minutes"),
      ("Automatic check-in", "people count ≥ 1 for one continuous minute"),
      ("Visible countdown", "the user can always claim the room manually"),
      ("LED strip", "Amber — booked, awaiting check-in")]),
]
for fn, title, sub, feats in panels:
    s = blank()
    header(s, "Panel experience", title, sub)
    s.shapes.add_picture(str(IMG / fn), Inches(0.72), Inches(2.1), width=Inches(7.5))
    for i, (h, t) in enumerate(feats):
        yy = 2.25 + i * 1.12
        rect(s, 8.55, yy, 4.06, 0.98, fill=RGBColor(0xFA, 0xFB, 0xFC), line=EDGE)
        tb(s, 8.8, yy + 0.15, 3.6, 0.28, h, 12.5, True, INK)
        tb(s, 8.8, yy + 0.47, 3.6, 0.45, t, 10.5, False, MUTE, line=1.25)
    footer(s, num() + 1)

# ---------------------------------------------------------------- 9 LED
s = blank()
header(s, "Availability indication", "The LED strip stays, under our control",
       "Switch the strip to manual once at setup, then drive it from the broker.")
mono(s, 0.72, 2.15, 11.89, 0.95,
     "xConfiguration UserInterface LedControl Mode: Manual        ← set once, part of Cisco's customer-managed setup\n"
     "xCommand UserInterface LedControl Color Set Color: Green|Yellow|Red|Off", 13)
states = [("Green", "Available", "No booking active, room free to take", GREEN, RGBColor(0xEC, 0xFA, 0xF1)),
          ("Amber", "Booked, awaiting check-in", "Within the check-in window, nobody detected yet", AMBER, RGBColor(0xFF, 0xF8, 0xEB)),
          ("Red", "In use", "Meeting running, or occupancy confirmed", RED, RGBColor(0xFE, 0xF3, 0xF3)),
          ("Off", "Out of service", "Panel unreachable or room disabled", MUTE, RGBColor(0xF4, 0xF6, 0xF8))]
yy = 3.35
for name, label, desc, col, bg in states:
    rect(s, 0.72, yy, 11.89, 0.72, fill=bg, line=None)
    rect(s, 0.92, yy + 0.16, 0.13, 0.4, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 1.3, yy + 0.22, 1.3, 0.3, name, 13.5, True, col)
    tb(s, 2.8, yy + 0.22, 3.4, 0.3, label, 13, True, INK)
    tb(s, 6.3, yy + 0.24, 6.1, 0.3, desc, 12, False, BODY)
    yy += 0.82
tb(s, 0.72, 6.75, 11.89, 0.3,
   "Cisco's own description of Auto mode is \u201cgreen: room available, red: room in use\u201d — we reproduce the same "
   "semantics, sourced from our calendar and sensor state instead of theirs.", 11.5, False, MUTE)
footer(s, num() + 1)

# ---------------------------------------------------------------- 10 booking flow
s = blank()
header(s, "Flow 1", "Ad-hoc booking from the panel",
       "User taps \u201cBook 15 min\u201d on a free room. Typical end-to-end time: under two seconds.")
steps = [("1", "Panel", "User taps a duration in the web app", DARK),
         ("2", "Broker", "Receives the request over the local network", BLUE),
         ("3", "Graph", "Creates the event with the broker as organizer\nand the room as a resource attendee", BLUE),
         ("4", "Exchange", "Resource booking attendant accepts or rejects —\nthis is what prevents double-booking", GREEN),
         ("5", "Broker", "Re-reads calendarView, updates panel state", BLUE),
         ("6", "Panel", "Shows \u201cIn use\u201d · LED set to Red", DARK)]
x = 0.72
for i, (n, who, what, col) in enumerate(steps):
    rect(s, x, 2.3, 1.82, 2.45)
    rect(s, x, 2.3, 1.82, 0.07, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x + 0.14, 2.48, 0.38, 0.38, fill=col, line=None, adj=0.5)
    tb(s, x + 0.14, 2.555, 0.38, 0.25, n, 12, True, WHITE, align=PP_ALIGN.CENTER)
    tb(s, x + 0.62, 2.56, 1.1, 0.25, who, 11.5, True, col)
    tb(s, x + 0.16, 3.05, 1.52, 1.6, what, 10.5, False, BODY, line=1.3, space=1)
    if i < 5:
        a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x + 1.86), Inches(3.35), Inches(0.21), Inches(0.18))
        a.fill.solid(); a.fill.fore_color.rgb = RGBColor(0xC3, 0xCC, 0xD6); a.line.fill.background(); a.shadow.inherit = False
    x += 2.07
rect(s, 0.72, 5.1, 11.89, 1.55, fill=RGBColor(0xFA, 0xFB, 0xFC), line=EDGE)
tb(s, 1.05, 5.3, 5.3, 0.3, "Why the broker is the organizer", 13.5, True, INK)
tb(s, 1.05, 5.68, 5.3, 0.85,
   "A room mailbox can only decline its own copy of a meeting. Only the organizer can truly cancel one. "
   "Owning the organizer identity is what lets us release a room cleanly later.", 11.5, False, BODY, line=1.3)
tb(s, 6.9, 5.3, 5.4, 0.3, "Why we do not write to the room calendar directly", 13.5, True, INK)
tb(s, 6.9, 5.68, 5.4, 0.85,
   "Writing straight into the room mailbox bypasses the resource booking attendant, so conflict rules are never "
   "evaluated and a room that must never double-book, silently does.", 11.5, False, BODY, line=1.3)
footer(s, num() + 1)

# ---------------------------------------------------------------- 11 release flow
s = blank()
header(s, "Flow 2", "Occupancy-driven room release",
       "The ghost-meeting problem: a room booked, nobody turns up, and it stays blocked all afternoon.")
tl = [("Meeting start", "Broker arms the check-in window and sets the LED amber", AMBER),
      ("0 – 10 min", "Panel shows \u201cCheck in\u201d. People count polled from the in-room video device", AMBER),
      ("Someone arrives", "People count ≥ 1 for one continuous minute → automatic check-in, LED red", GREEN),
      ("Nobody arrives", "30-second countdown shown on the panel, then the room is released", RED),
      ("Release", "Broker-owned booking is cancelled outright. User-organised booking has the room's copy declined", RED),
      ("Room freed", "LED back to green, slot immediately bookable by anyone", GREEN)]
yy = 2.18
for i, (t, d, col) in enumerate(tl):
    rect(s, 1.25, yy + 0.1, 0.16, 0.6 if i < 5 else 0.0, fill=RGBColor(0xE4, 0xE9, 0xEF), line=None, shape=MSO_SHAPE.RECTANGLE)
    rect(s, 1.13, yy + 0.08, 0.4, 0.4, fill=col, line=None, adj=0.5)
    tb(s, 1.75, yy + 0.06, 2.9, 0.3, t, 13.5, True, INK)
    tb(s, 4.75, yy + 0.09, 7.8, 0.3, d, 12, False, BODY)
    yy += 0.68
rect(s, 0.72, 6.18, 11.89, 0.7, fill=RGBColor(0xFF, 0xF8, 0xEB), line=RGBColor(0xF5, 0xD9, 0x9C))
tb(s, 1.05, 6.32, 11.3, 0.45,
   "Honest limitation: for a meeting organised by a user in Outlook we can free the room, but we cannot delete their "
   "meeting. They get a decline notice and the entry stays on their own calendar.", 11.5, False,
   RGBColor(0x92, 0x5F, 0x08), line=1.25)
footer(s, num() + 1)

# ---------------------------------------------------------------- 12 graph
s = blank()
header(s, "Integration", "Microsoft Graph — Exchange Online",
       "App-only daemon, scoped so it can only ever touch room mailboxes.")
blocks = [
    ("Identity & scope", BLUE,
     ["Entra app, client-credentials (app-only) flow",
      "Place.Read.All — discover rooms",
      "Calendars.ReadWrite — read and write bookings",
      "Scoped with RBAC for Applications, so the app reaches only the room mailboxes, never the wider tenant"]),
    ("Reading", GREEN,
     ["calendarView — occurrences already expanded, including modified recurring instances",
      "calendarView/delta for incremental sync",
      "Polling every 1–5 min, so no public endpoint and no inbound firewall exposure"]),
    ("Writing", AMBER,
     ["Create as organizer, room as resource attendee",
      "Exchange arbitrates conflicts natively",
      "Cancel our own bookings outright",
      "Decline the room's copy for user-organised ones"]),
]
x = 0.72
for t, c, items in blocks:
    rect(s, x, 2.2, 3.86, 3.55)
    rect(s, x, 2.2, 3.86, 0.07, fill=c, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(s, x + 0.28, 2.48, 3.3, 0.3, t, 15, True, INK)
    tb(s, x + 0.28, 2.98, 3.32, 2.6, ["·  " + it for it in items], 11.5, False, BODY, line=1.28, space=10)
    x += 4.03
rect(s, 0.72, 5.95, 11.89, 0.78, fill=RGBColor(0xF1, 0xF6, 0xFE), line=RGBColor(0xC7, 0xDC, 0xFC))
tb(s, 1.05, 6.1, 11.3, 0.5,
   "Better privacy posture than before: Webex Hybrid Calendar required tenant-wide \u201cread/write calendars in all "
   "mailboxes\u201d consent. This design is restricted to the room mailboxes alone.", 13, False, RGBColor(0x1E, 0x40, 0xAF), line=1.25)
footer(s, num() + 1)

# ---------------------------------------------------------------- 13 admin UI
s = blank()
header(s, "Management", "Pairing a panel, a video device and a mailbox",
       "One screen defines a room: which panel, which sensor source, which calendar.")
s.shapes.add_picture(str(IMG / "admin_pair.png"), Inches(2.67), Inches(2.02), width=Inches(8.0))
footer(s, num() + 1)

# ---------------------------------------------------------------- 14 admin detail
s = blank()
header(s, "Management", "What pairing actually configures",
       "The broker validates each device on save, so misconfiguration is caught at setup rather than at 09:00 Monday.")
items = [("Booking panel", "Verifies Provisioning Mode: Off, switches the panel to Persistent Web App mode, pushes the "
                           "web app URL, and sets LedControl Mode to Manual.", DARK),
         ("Video device", "Verifies RoomAnalytics PeopleCountOutOfCall is On, subscribes to people count and presence "
                          "over the local xAPI. Optional — without it, release falls back to check-in only.", BLUE),
         ("Room mailbox", "Picked from rooms discovered through the Graph Places API. Confirms the token is valid and "
                          "scoped to that mailbox before the room goes live.", GREEN),
         ("Policy", "Check-in window, release-if-empty threshold, and how much meeting detail the panel is allowed "
                    "to display in a public corridor.", AMBER)]
yy = 2.2
for t, d, c in items:
    rect(s, 0.72, yy, 11.89, 1.08)
    rect(s, 0.72, yy, 0.08, 1.08, fill=c, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 1.15, yy + 0.2, 2.5, 0.3, t, 14, True, INK)
    tb(s, 3.85, yy + 0.2, 8.5, 0.7, d, 12, False, BODY, line=1.3)
    yy += 1.2
footer(s, num() + 1)

# ---------------------------------------------------------------- 15 features
s = blank()
header(s, "Scope", "Feature summary", "What the concept delivers against what Webex did before.")
cols = [("Delivered", GREEN, RGBColor(0xF1, 0xF9, 0xF4),
         ["Room status on a wall panel, with LED strip",
          "Ad-hoc booking from the panel",
          "Today's agenda from the room mailbox",
          "Extend and end-early from the panel",
          "Check-in and automatic room release",
          "Occupancy from the in-room video device",
          "Temperature and humidity on the panel",
          "Multi-calendar: Graph, EWS, Google, CalDAV",
          "Central management and pairing UI"]),
        ("Lost with Webex", RED, RGBColor(0xFE, 0xF4, 0xF4),
         ["Cisco's native Room Scheduler screen",
          "One Button To Push on the panel itself",
          "Control Hub device management and firmware",
          "Webex cloud analytics and workspace insights",
          "Cisco TAC support for the booking experience",
          "@webex / @meet keyword meeting rewriting"])]
x = 0.72
for t, c, bg, items in cols:
    w = 6.3 if t == "Delivered" else 5.4
    rect(s, x, 2.2, w, 4.45, fill=bg, line=None)
    tb(s, x + 0.35, 2.45, w - 0.7, 0.35, t.upper(), 11.5, True, c)
    yy = 2.95
    for it in items:
        tb(s, x + 0.35, yy, w - 0.7, 0.3, ("✓   " if c == GREEN else "×   ") + it, 12.5, False,
           RGBColor(0x14, 0x6C, 0x43) if c == GREEN else RGBColor(0x99, 0x1B, 0x1B))
        yy += 0.4
    x += w + 0.45
footer(s, num() + 1)

# ---------------------------------------------------------------- 16 risks
s = blank()
header(s, "Before committing", "Open items and honest risks",
       "Nothing here is a blocker, but each needs a decision or a lab test.")
risks = [("Must be tested on real hardware", RED,
          "Cisco's stand-alone guide dates from April 2024 and has never been revised. The current xAPI schema lists "
          "the Bookings family for the Navigator without a customer-managed caveat. Worth one lab test: does Scheduler "
          "mode work unregistered? If it does, the native UI comes back on the table."),
         ("Ghost meetings on organiser calendars", AMBER,
          "Releasing a user-organised meeting frees the room but leaves the entry on their calendar. A UX and comms "
          "decision, not a technical one."),
         ("Sensor accuracy is unpublished", AMBER,
          "Cisco publishes no figures for people-count latency or false negatives. Characterise in one pilot room "
          "before setting an aggressive release threshold."),
         ("We own what Cisco used to own", BLUE,
          "Firmware updates, certificates, panel provisioning and support all become ours. This is the real cost of "
          "leaving the cloud, and it is ongoing rather than one-off.")]
yy = 2.2
for t, c, d in risks:
    rect(s, 0.72, yy, 11.89, 1.07)
    rect(s, 0.72, yy, 0.08, 1.07, fill=c, line=None, shape=MSO_SHAPE.RECTANGLE)
    tb(s, 1.15, yy + 0.17, 4.1, 0.5, t, 13.5, True, INK, line=1.15)
    tb(s, 5.5, yy + 0.16, 6.9, 0.8, d, 11, False, BODY, line=1.25)
    yy += 1.17
footer(s, num() + 1)

# ---------------------------------------------------------------- 17 next
s = blank()
header(s, "Next steps", "Proving it in stages", "Each stage is independently useful and independently abandonable.")
ph = [("1", "Spike", "One Navigator, customer managed, Persistent Web App.\nProve the panel renders our UI and the LED "
                     "responds to\nLedControl Color Set.", "1 day"),
      ("2", "Read-only", "Graph driver against one room mailbox. Panel shows real\nbookings. No writes anywhere.", "1 week"),
      ("3", "Booking", "Ad-hoc booking from the panel, broker as organizer.\nConflict handling via Exchange.", "2 weeks"),
      ("4", "Release", "Subscribe to the video device, occupancy state machine,\ncheck-in and auto-release.", "2 weeks"),
      ("5", "Scale", "Management UI, additional calendar drivers,\nmonitoring and alerting across all rooms.", "ongoing")]
x = 0.72
for n, t, d, dur in ph:
    rect(s, x, 2.3, 2.26, 3.3)
    rect(s, x, 2.3, 2.26, 0.07, fill=BLUE, line=None, shape=MSO_SHAPE.RECTANGLE)
    rect(s, x + 0.2, 2.52, 0.42, 0.42, fill=BLUE, line=None, adj=0.5)
    tb(s, x + 0.2, 2.615, 0.42, 0.28, n, 13, True, WHITE, align=PP_ALIGN.CENTER)
    tb(s, x + 0.2, 3.1, 1.9, 0.3, t, 14.5, True, INK)
    tb(s, x + 0.2, 3.5, 1.95, 1.6, d, 10, False, BODY, line=1.3)
    chip(s, x + 0.2, 5.05, 1.1, 0.33, dur, MUTE, RGBColor(0xF0, 0xF3, 0xF6))
    x += 2.42
rect(s, 0.72, 5.95, 11.89, 0.85, fill=RGBColor(0x10, 0x28, 0x1A), line=None)
tb(s, 1.05, 6.14, 11.3, 0.5,
   "Recommended first move: stage 1. One panel, one afternoon, and it settles the only genuinely open "
   "technical question in this deck.", 14, True, RGBColor(0x4A, 0xDE, 0x80))
footer(s, num() + 1)

out = D / "Self-hosted-room-booking-concept.pptx"
prs.save(str(out))
print("saved", out, out.stat().st_size, "bytes,", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
